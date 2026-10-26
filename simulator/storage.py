from config import EXCEL_FILE

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    Workbook = None
    load_workbook = None

def create_excel(self):
    if Workbook is None:
        return

    try:
        if EXCEL_FILE.exists():
            wb = load_workbook(EXCEL_FILE)
            if "Объекты" in wb.sheetnames:
                ws = wb["Объекты"]
            else:
                ws = wb.create_sheet("Объекты")
                ws.append(["ID", "Имя", "Модель", "Маршрут"])

            # Гарантируем заголовки.
            if ws.max_row == 1 and ws["A1"].value is None:
                ws.append(["ID", "Имя", "Модель", "Маршрут"])

            # Записываем только текст маршрута и основные данные.
            # Координаты и точки траектории в Excel НЕ записываются.
            for d in self.drones:
            # Последовательно обрабатываем каждый элемент коллекции.
                row = d.id + 1
                ws.cell(row=row, column=1, value=d.id)
                ws.cell(row=row, column=2, value=d.name)
                ws.cell(row=row, column=3, value=d.model)
                ws.cell(row=row, column=4, value=d.route_text)

            wb.save(EXCEL_FILE)
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Объекты"
            ws.append(["ID", "Имя", "Модель", "Маршрут"])

            for d in self.drones:
                ws.append([d.id, d.name, d.model, d.route_text])

            wb.save(EXCEL_FILE)

    except Exception as e:
    # Обрабатываем ошибку и не даём ей остановить работу всей симуляции.
        print("Не удалось обновить Excel:", e)
